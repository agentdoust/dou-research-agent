"""
Agente de IA para Pesquisa no Diário Oficial da União (DOU)

Persona: Jornalista isento e técnico, responsável pelas pesquisas oriundas de fontes oficiais.

Fonte oficial: https://pesquisa.in.gov.br/imprensa/core/jornalList.action

Tarefa:
  - Seção 1 (DOU1): publicações de atos normativos relacionados à segurança pública,
    restritas ao Ministério da Justiça e Segurança Pública.
  - Seção 2 (DOU2): nomeações e exonerações relacionadas à Polícia Federal,
    restritas ao Ministério da Justiça e Segurança Pública.

Saída: documento Word gerado na pasta Resultado (no local usa D:\\opencode_desktop\\resultado).
"""

import argparse
import re
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

import pymupdf
import requests
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------
# Caminhos portáveis: em máquina local usa-se D:\opencode_desktop; no deploy
# (Streamlit Cloud/Linux) o filesystem é somente leitura exceto /tmp, então os
# diretórios de download e de resultados caem para o tempdir do usuário.
AGENT_DIR = Path(__file__).resolve().parent
HOST_BASE = Path(r"D:\opencode_desktop")
RAIZ_REPO = AGENT_DIR.parent.parent if AGENT_DIR.parent.name == "Agents" else AGENT_DIR.parent


def _escolher_diretorio(nome, *candidatos):
    """Retorna o primeiro diretório gravável entre os candidatos; senão usa tempdir."""
    for base in candidatos:
        try:
            d = base / nome
            d.mkdir(parents=True, exist_ok=True)
            (d / ".gravavel").write_text("ok")
            (d / ".gravavel").unlink()
            return d
        except OSError:
            continue
    d = Path(tempfile.gettempdir()) / "dou_agent" / nome
    d.mkdir(parents=True, exist_ok=True)
    return d


RESULT_DIR = _escolher_diretorio("resultado", HOST_BASE, RAIZ_REPO)
CACHE_DIR = _escolher_diretorio("downloads", AGENT_DIR, RAIZ_REPO)

PESQUISA_URL = "https://pesquisa.in.gov.br/imprensa/core/jornalList.action"
START_URL = "https://pesquisa.in.gov.br/imprensa/core/start.action"

# Sequência principal de jornais por seção (DOU1 = 515, DOU2 = 529)
SECOES = {
    1: {"jornal": "515", "arquivo": "ASSINADO_do1", "titulo": "Diário Oficial da União - Seção 1"},
    2: {"jornal": "529", "arquivo": "ASSINADO_do2", "titulo": "Diário Oficial da União - Seção 2"},
}

MINISTERIO_ALVO = "Ministério da Justiça e Segurança Pública"

# Linhas que aparecem nos cabeçalhos/rodapés de página e não são conteúdo
NOISE_RE = re.compile(
    r"Documento assinado digitalmente conforme MP|"
    r"que institui a Infraestrutura de Chaves Públicas|"
    r"Este documento pode ser verificado|"
    r"http://www\.in\.gov\.br/autenticidade|"
    r"^ISSN\s|"
    r"^N[º°]\s+\d+\s*,\s*(segunda|terça|quarta|quinta|sexta|sábado|domingo)", re.I
)

# Cabeçalho de ministério: linha curta com nome de órgão, sem números
# (evita confundir texto contínuo tipo "Ministério da Justiça..., publicada no DOU nº ...")
MIN_HEADER_RE = re.compile(r"^Minist[eé]rio [^\d\n]{1,90}$")
MJSP_RE = re.compile(r"^Minist[eé]rio da Justi[çc]a e Seguran[çc]a P[úu]blica$")

ACT_HEADER_RE = re.compile(
    r"^\s*(ALVAR[ÁA] N[º°]\s*\d+|"
    r"PORTARIA(S)? [A-Z0-9º.\- ]{2,90}N[º°]\s*\d+|"
    r"PORTARIA [A-Z0-9º.\- ]{2,40}DE \d{1,2} DE \w+ DE \d{4}|"
    r"RESOLU[ÇC][AÃ]O|"
    r"INSTRU[ÇC][AÃ]O NORMATIVA.*N[º°]\s*\d+|"
    r"DECRETO.*N[º°]\s*\d+|"
    r"EDITAL.*N[º°]\s*\d+|"
    r"RETIFICA[ÇC][AÃ]O|"
    r"DESPACHO[^\n]*\d|"
    r"AVISO[^\n]*\d)"
)

ASSINADO_RE = re.compile(r"redirecionaSelect\('(https://download\.in\.gov\.br/[^']*ASSINADO[^']*)'\);")

UA = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    )
}


# ---------------------------------------------------------------------------
# Funções de rede / download
# ---------------------------------------------------------------------------
def criar_sessao() -> requests.Session:
    s = requests.Session()
    s.headers.update(UA)
    s.get(START_URL, timeout=30)  # estabelece cookies/sessão
    return s


def obter_link_edicao(sessao, secao: int, data: date):
    cfg = SECOES[secao]
    params = {
        "edicao.dtInicio": f"{data.day:02d}/{data.month:02d}",
        "edicao.dtFim": f"{data.day:02d}/{data.month:02d}",
        "edicao.ano": str(data.year),
        "edicao.jornal": cfg["jornal"],
    }
    r = sessao.get(PESQUISA_URL, params=params, timeout=60)
    r.raise_for_status()
    m = ASSINADO_RE.search(r.text)
    if not m:
        return None
    return m.group(1).replace("&amp;", "&")


def baixar_pdf(url: str, destino: Path):
    destino.parent.mkdir(parents=True, exist_ok=True)
    if destino.exists() and destino.stat().st_size > 1000:
        print(f"  [cache] {destino.name} já existe ({destino.stat().st_size // 1024} KB)")
        return True
    print(f"  [download] {url[:110]}...")
    with requests.get(url, headers=UA, timeout=900, stream=True) as r:
        r.raise_for_status()
        with open(destino, "wb") as f:
            for chunk in r.iter_content(1 << 20):
                f.write(chunk)
    print(f"  [ok] {destino.stat().st_size // 1024} KB")
    return True


def ler_paginas(pdf_path: Path):
    doc = pymupdf.open(str(pdf_path))
    try:
        return [p.get_text() for p in doc]
    finally:
        doc.close()


def extrair_texto(pdf_path: Path) -> str:
    return "\n\f".join(ler_paginas(pdf_path))


# ---------------------------------------------------------------------------
# Índices (Sumário) das seções do DOU
# ---------------------------------------------------------------------------
# Linhas do sumário na 1ª página: "Ministério da Justiça e Segurança Pública .... 81"
SUMARIO_RE = re.compile(r"^\s*(?P<nome>.+?)\s*[\.\s]{4,}\s*(?P<pagina>\d{1,4})\s*$")
FIM_SUMARIO_RE = re.compile(r"[\.\s]{4,}Esta edi[çc][ãa]o [eé] composta de \d+ p[áa]ginas")


def ler_indices(paginas):
    """Extrai, em ordem, a lista (nome, página inicial) do sumário da seção."""
    indices = []
    for i, pag in enumerate(paginas[:2]):
        for ln in pag.split("\n"):
            l = ln.rstrip()
            if FIM_SUMARIO_RE.search(l):
                return indices
            m = SUMARIO_RE.match(l)
            if m:
                nome = m.group("nome").strip()
                if nome and not nome.startswith("."):
                    indices.append({
                        "nome": nome,
                        "pagina": int(m.group("pagina")),
                    })
    return indices


def _inicios_paginas(paginas):
    """Índice, no texto unido, da 1ª linha de cada página do PDF."""
    inicios, pos = [], 0
    for k, pag in enumerate(paginas):
        inicios.append(pos)
        pos += len(pag.split("\n")) + (0 if k == len(paginas) - 1 else 1)
    return inicios


def _linha_cabecalho_indice(linhas, inicios, indice, pagina_pdf):
    """Linha do cabeçalho exato do índice no conteúdo (após a página do sumário).

    O número de página do sumário pode divergir ±1 da página física do PDF;
    a varredura começa na página anterior à indicada e cobre até 2 páginas
    adiante, tolerando a diferença sem varrer o documento inteiro.
    """
    ini = inicios[max(0, pagina_pdf - 1)]
    fim = inicios[min(pagina_pdf + 2, len(inicios) - 1)]
    alvo = indice["nome"].casefold()
    for i in range(ini, min(fim, len(linhas))):
        if linhas[i].strip().casefold() == alvo:
            return i
    return ini


def fatiar_indices(texto: str, paginas, nomes_selecionados):
    """Limita o texto aos blocos dos índices selecionados (sumário da edição).

    Bloco de cada índice = do cabeçalho do índice até o cabeçalho do índice
    seguinte (ou fim do documento).
    """
    linhas = texto.split("\n")
    inicios = _inicios_paginas(paginas)
    todos = ler_indices(paginas)
    escolhidos = [t for t in todos if t["nome"] in nomes_selecionados]
    if not escolhidos:
        return texto

    # cabeçalho de cada índice no conteúdo (descartando a página do sumário)
    cab = {}
    for t in todos:
        if t["pagina"] - 1 >= len(paginas):
            cab[t["nome"]] = None
            continue
        cab[t["nome"]] = _linha_cabecalho_indice(linhas, inicios, t, t["pagina"] - 1)

    pedacos = []
    for k, t in enumerate(escolhidos):
        ini = cab[t["nome"]]
        if ini is None:
            continue
        prox = next(
            (
                u for u in todos
                if (u["pagina"], todos.index(u)) > (t["pagina"], todos.index(t))
            ),
            None,
        )
        fim = cab[prox["nome"]] if prox and cab.get(prox["nome"]) is not None else len(linhas)
        pedacos.append("\n".join(linhas[ini:fim]))
    if not pedacos:
        return texto
    return "\n\f".join(pedacos)


# ---------------------------------------------------------------------------
# Segmentação por órgão (Ministério da Justiça e Segurança Pública)
# ---------------------------------------------------------------------------
def achar_bloco_mjsp(texto: str):
    linhas = texto.split("\n")
    cab = [i for i, ln in enumerate(linhas) if MIN_HEADER_RE.match(ln.strip())]
    inicio = next((i for i in cab if MJSP_RE.match(linhas[i].strip())), None)
    if inicio is None:
        return None
    fim = next((j for j in cab if j > inicio), len(linhas))
    return inicio, fim


def limpar_linha(linha: str) -> bool:
    l = linha.strip()
    if not l:
        return True
    if NOISE_RE.search(l):
        return True
    if re.match(r"^\d{1,3}$", l):
        return True
    return False


# ---------------------------------------------------------------------------
# Extração da Seção 1 — atos normativos (segurança pública) do MJSP
# ---------------------------------------------------------------------------
def extrair_atos_secao1(linhas):
    atos = []
    i = 0
    while i < len(linhas):
        l = linhas[i].strip()
        if ACT_HEADER_RE.match(l):
            titulo = l
            j = i + 1
            corpo = []
            while j < len(linhas):
                lj = linhas[j].strip()
                if ACT_HEADER_RE.match(lj):
                    break
                if not limpar_linha(linhas[j]):
                    corpo.append(lj)
                j += 1
            ato = {
                "titulo": titulo,
                "corpo": " ".join(corpo),
                "fim": j,
            }
            if eh_ato_seguranca_publica(ato):
                atos.append(ato)
            i = j
        else:
            i += 1
    return atos


def resumir_corpo(corpo: str, max_len: int = 900) -> str:
    corpo = re.sub(r"\s+", " ", corpo).strip()
    idx = corpo.lower().find("resolve:")
    texto = corpo[idx + len("resolve:"):] if idx != -1 else corpo
    texto = texto.strip(" .;,:-")
    if len(texto) > max_len:
        texto = texto[:max_len].rstrip() + "..."
    return texto


# Termos que caracterizam atos de segurança pública (Seção 1).
# Específicos do tema: bastam sozinhos para classificar um ato.
SEGURANCA_TERMOS = [
    "alvará", "munição", "munições", "pólvora", "espoleta", "revólver",
    "espingarda", "pistola", "calibre", "entorpecente", "penitenciária",
    "sistema prisional", "polícia penal", "tráfico de", "latrocínio",
    "estelionato", "sequestro", "criminalidade", "segurança penitenciária",
]


def eh_ato_seguranca_publica(ato) -> bool:
    texto = (ato["titulo"] + " " + ato["corpo"]).lower()
    return any(t in texto for t in SEGURANCA_TERMOS)


# ---------------------------------------------------------------------------
# Extração da Seção 2 — nomeações/exonerações Polícia Federal (MJSP)
# ---------------------------------------------------------------------------
def extrair_nomeacoes_secao2(linhas):
    """Localiza o sub-bloco 'POLÍCIA FEDERAL' do MJSP e extrai itens de pessoal."""
    # 1) delimita o sub-bloco PF dentro do bloco MJSP
    ini = next((i for i, l in enumerate(linhas) if re.match(r"^POL[ÍI]CIA FEDERAL$", l.strip().upper())), None)
    if ini is None:
        return []

    def saiu_pf(l):
        lu = l.strip().upper()
        if MIN_HEADER_RE.match(lu):  # cabeçalho de outro ministério
            return True
        return re.match(
            r"^(SECRETARIA NACIONAL|AG[ÊE]NCIA NACIONAL|CONSELHO ADMINISTRATIVO|"
            r"POL[ÍI]CIA RODOV[ÍI]ARIA FEDERAL|DEPARTAMENTO PENITENCI[ÁA]RIO NACIONAL|"
            r"PORTARIA DE PESSOAL|PORTARIA SENAD|PORTARIA GABPR|PORTARIA CADE)",
            lu,
        )

    fim = next((i for i in range(ini + 1, len(linhas)) if saiu_pf(linhas[i])), len(linhas))
    bloco = linhas[ini + 1:fim]

    # 2) separa o bloco PF em atos (portarias) e divide itens de pessoal
    atos = []
    corrente = None
    for l in bloco:
        lu = l.strip().upper()
        if re.match(r"^PORTARIA(S)? (DG/PF|DGP|DDG|DGP/PF)", lu):
            if corrente is not None:
                atos.append(corrente)
            corrente = {"titulo": l.strip(), "linhas": []}
        elif corrente is not None:
            corrente["linhas"].append(l)
    if corrente is not None:
        atos.append(corrente)

    resultado = []
    for ato in atos:
        juntado = re.sub(r"\s+", " ", " ".join(ato["linhas"]))
        # localiza cada item de pessoal (número opcional + verbo) no texto contínuo
        itens = list(re.finditer(
            r"(?!\s)(?:N[º°]\s*\d+[A-Z]?\s+)?(?:Nomear|Designar|Exonerar|Dispensar)\b",
            juntado,
            flags=re.IGNORECASE,
        ))
        spans = [it.start() for it in itens]
        segs = [juntado[s:e].strip() for s, e in zip(
            spans, spans[1:] + [len(juntado)]
        )]
        for seg in segs:
            mnum = re.match(r"(?P<num>N[º°]\s*\d+[A-Z]?)\s+", seg)
            fim_num = mnum.end() if mnum else 0
            mverbo = re.match(
                r"(?P<verbo>Nomear|Designar|Exonerar|Dispensar)\b",
                seg[fim_num:],
                flags=re.IGNORECASE,
            )
            if not mverbo:
                continue
            depois = seg[mverbo.end():]
            mnome = re.search(r"[A-ZÀ-Ú]{2,}(?:\s+[A-ZÀ-Ú]{2,}){1,7}", depois)
            if not mnome:
                continue
            nome = mnome.group(0)
            resto = _limpar_resto(depois[mnome.end():])
            resultado.append({
                "ato": ato["titulo"],
                "num": (mnum.group("num") if mnum else "").strip(),
                "verbo": mverbo.group("verbo").capitalize(),
                "nome": nome,
                "resto": resto,
            })
    return resultado


def _limpar_resto(texto):
    """Remove assinaturas/páginas residuais e normaliza o trecho de detalhes."""
    texto = re.sub(r"\s+", " ", texto).strip()
    texto = re.sub(
        r",?\s*(GESTÃO DE PESSOAS|HELENA DE REZENDE|ANDREI AUGUSTO PASSOS RODRIGUES)?$",
        "",
        texto,
    )
    texto = texto.strip(" ,.")
    return texto


# ---------------------------------------------------------------------------
# Geração do documento Word
# ---------------------------------------------------------------------------
def gerar_docx(data, resultados):
    RESULT_DIR.mkdir(exist_ok=True)
    nome = f"Relatorio_DOU_MJSP_{data.strftime('%Y_%m_%d')}.docx"
    caminho = RESULT_DIR / nome

    doc = Document()
    titulo = doc.add_heading("Relatório de Pesquisa — Diário Oficial da União", 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Fonte: {PESQUISA_URL}\n").italic = True
    p.add_run(f"Edição de {data.strftime('%d/%m/%Y')} — filtro: {MINISTERIO_ALVO}\n").italic = True
    p.add_run(f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}").italic = True

    # Seção 1
    doc.add_heading("Seção 1 — Atos normativos de segurança pública (MJSP)", level=1)
    s1 = resultados.get("secao1", [])
    doc.add_paragraph(f"Total de atos identificados: {len(s1)}")
    if s1:
        for k, a in enumerate(s1, 1):
            hp = doc.add_paragraph()
            run = hp.add_run(f"{k}. {a['titulo']}")
            run.bold = True
            corpo = resumir_corpo(a["corpo"])
            if corpo:
                doc.add_paragraph(corpo, style="List Bullet")
    else:
        doc.add_paragraph("Nenhum ato normativo de segurança pública identificado no MJSP.")

    # Seção 2
    doc.add_heading(
        "Seção 2 — Pessoal da Polícia Federal (MJSP): nomeações, exonerações, "
        "designações e dispensas",
        level=1,
    )
    s2 = resultados.get("secao2", [])
    doc.add_paragraph(f"Total de ocorrências identificadas: {len(s2)}")
    if s2:
        tabela = doc.add_table(rows=1, cols=5)
        tabela.style = "Light Grid Accent 1"
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = tabela.rows[0].cells
        for c, txt in enumerate(["Nº", "Ato", "Ação", "Nome", "Detalhes"]):
            hdr[c].text = txt
            for r_ in hdr[c].paragraphs[0].runs:
                r_.bold = True
        for item in s2:
            cells = tabela.add_row().cells
            cells[0].text = item["num"]
            cells[1].text = item["ato"][:45]
            cells[2].text = item["verbo"]
            cells[3].text = item["nome"]
            cells[4].text = item["resto"][:200]
        for w, width in zip(range(5), [Inches(0.5), Inches(1.8), Inches(1.0), Inches(1.6), Inches(2.2)]):
            for row in tabela.rows:
                row.cells[w].width = width
    else:
        doc.add_paragraph(
            "Nenhuma ocorrência de pessoal da Polícia Federal identificada no MJSP "
            "(nomeações, exonerações, designações e dispensas)."
        )

    doc.save(str(caminho))
    return caminho


# ---------------------------------------------------------------------------
# Fluxo principal
# ---------------------------------------------------------------------------
def processar_secao(sessao, secao: int, data: date, indices=None) -> dict:
    cfg = SECOES[secao]
    print(f"\n--- Seção {secao} ({cfg['titulo']}) ---")
    link = obter_link_edicao(sessao, secao, data)
    if not link:
        print("  [aviso] Edição completa (ASSINADO) não encontrada no jornalList.action")
        return {"secao1": []} if secao == 1 else {"secao2": []}

    arquivo = link.split("/")[-1].split("?")[0]
    destino = CACHE_DIR / data.strftime("%Y_%m_%d") / arquivo
    baixar_pdf(link, destino)

    paginas = ler_paginas(destino)
    texto = "\n\f".join(paginas)

    if indices:
        texto = fatiar_indices(texto, paginas, indices)
        print(f"  [ok] Índices selecionados: {', '.join(indices)}")

    bloco = achar_bloco_mjsp(texto)
    if not bloco:
        print(f"  [aviso] Bloco '{MINISTERIO_ALVO}' não localizado")
        return {"secao1": []} if secao == 1 else {"secao2": []}

    linhas = texto.split("\n")[bloco[0]:bloco[1]]
    print(f"  [ok] Bloco MJSP: linhas {bloco[0]+1}..{bloco[1]} ({len(linhas)} linhas)")

    if secao == 1:
        atos = extrair_atos_secao1(linhas)
        print(f"  [ok] {len(atos)} atos normativos localizados")
        return {"secao1": atos}
    else:
        nomes = extrair_nomeacoes_secao2(linhas)
        print(f"  [ok] {len(nomes)} ocorrências de pessoal PF localizadas (nomeações, exonerações, designações e dispensas)")
        return {"secao2": nomes}


def main():
    parser = argparse.ArgumentParser(description="Agente de pesquisa no DOU (fonte oficial)")
    parser.add_argument(
        "--data",
        default=date.today().strftime("%d/%m/%Y"),
        help="Data da edição no formato DD/MM/AAAA (padrão: hoje)",
    )
    parser.add_argument(
        "--secoes",
        default="1,2",
        help="Seções a processar, separadas por vírgula (padrão: 1,2)",
    )
    parser.add_argument(
        "--indices",
        default="",
        help=(
            "Índices (entradas do sumário do DOU) a pesquisar, separados por ';'. "
            "Ex.: --indices \"Ministério da Justiça e Segurança Pública\". "
            "Sem limite de seção: aplica-se a todas as seções processadas."
        ),
    )
    args = parser.parse_args()

    try:
        data = datetime.strptime(args.data, "%d/%m/%Y").date()
    except ValueError:
        print(f"Data inválida: {args.data}. Use o formato DD/MM/AAAA.")
        sys.exit(2)

    indices_arg = [i.strip() for i in args.indices.split(";") if i.strip()] or None

    print("=" * 70)
    print("Agente de Pesquisa no Diário Oficial da União (DOU) — fonte oficial")
    print(f"Data da edição: {data.strftime('%d/%m/%Y')}")
    print(f"Órgão alvo: {MINISTERIO_ALVO}")
    print("=" * 70)

    sessao = criar_sessao()
    resultados = {}
    for s in map(int, args.secoes.split(",")):
        res = processar_secao(sessao, s, data, indices=indices_arg)
        resultados.update(res)

    print("\nGerando relatório Word...")
    caminho = gerar_docx(data, resultados)
    print(f"Relatório gerado: {caminho}")
    print("Concluído.")


if __name__ == "__main__":
    main()